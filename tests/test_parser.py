"""Tests for PDF/DOCX parsing error handling and edge cases."""

import io
import pytest
from analyzer.parser import (
    extract_text_from_file,
    extract_text_from_pdf,
    extract_text_from_docx,
    clean_extracted_text,
    ResumeParseResult
)

# Test pymupdf availability
try:
    import pymupdf as fitz
except ImportError:
    import fitz

import docx


def test_clean_extracted_text():
    raw = "  Hello \xa0 World \r\n\r\n\n\n\n Python Developer   "
    cleaned = clean_extracted_text(raw)
    assert "Hello   World" in cleaned
    assert "Python Developer" in cleaned
    assert "\n\n\n" not in cleaned


def test_empty_pdf_bytes():
    result = extract_text_from_pdf(b"")
    assert not result.success
    assert "empty" in result.error_message.lower()


def test_corrupted_pdf_bytes():
    result = extract_text_from_pdf(b"this is not a valid pdf at all %PDF-invalid")
    assert not result.success
    assert "could not open pdf" in result.error_message.lower() or "corrupted" in result.error_message.lower()


def test_valid_pdf_generation_and_parsing():
    # Create real in-memory PDF using fitz
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "John Doe\nPython Developer\nExperienced in Django, FastAPI, SQL, and Docker.")
    pdf_bytes = doc.tobytes()
    doc.close()

    result = extract_text_from_file(pdf_bytes, "resume.pdf")
    assert result.success
    assert result.file_type == "pdf"
    assert result.page_count == 1
    assert "Python Developer" in result.text
    assert "Django" in result.text
    assert not result.is_scanned


def test_scanned_pdf_warning():
    # Empty page with 0 text
    doc = fitz.open()
    doc.new_page()  # Blank page
    pdf_bytes = doc.tobytes()
    doc.close()

    result = extract_text_from_pdf(pdf_bytes)
    assert result.success
    assert result.is_scanned
    assert result.warning_message is not None
    assert "scanned" in result.warning_message.lower()


def test_empty_docx_bytes():
    result = extract_text_from_docx(b"")
    assert not result.success
    assert "empty" in result.error_message.lower()


def test_corrupted_docx_bytes():
    result = extract_text_from_docx(b"corrupted binary data not a zip file")
    assert not result.success
    assert "could not read docx" in result.error_message.lower()


def test_valid_docx_generation_and_parsing():
    # Create real in-memory DOCX
    doc = docx.Document()
    doc.add_heading("Jane Smith - AI Engineer", level=1)
    doc.add_paragraph("Skills: PyTorch, Machine Learning, Python, NLP, and Git.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Education"
    table.rows[0].cells[1].text = "B.Tech in Computer Science"

    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    result = extract_text_from_file(docx_bytes, "candidate.docx")
    assert result.success
    assert result.file_type == "docx"
    assert "Jane Smith" in result.text
    assert "PyTorch" in result.text
    assert "Education" in result.text


def test_unsupported_file_extension():
    result = extract_text_from_file(b"some data", "resume.exe")
    assert not result.success
    assert "unsupported file format" in result.error_message.lower()
