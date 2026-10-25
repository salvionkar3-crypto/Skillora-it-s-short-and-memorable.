"""Resume document parser supporting PDF and DOCX formats safely."""

import io
import re
from dataclasses import dataclass
from typing import Optional, Union, BinaryIO

# Try importing pymupdf (fitz)
try:
    import pymupdf as fitz
except ImportError:
    try:
        import fitz
    except ImportError:
        fitz = None

# Try importing docx
try:
    import docx
except ImportError:
    docx = None


@dataclass
class ResumeParseResult:
    """Structured result returned by resume text extraction."""
    success: bool
    text: str = ""
    file_type: str = "unknown"
    page_count: int = 0
    word_count: int = 0
    is_scanned: bool = False
    error_message: Optional[str] = None
    warning_message: Optional[str] = None


def clean_extracted_text(raw_text: str) -> str:
    """Clean and normalize extracted resume text."""
    if not raw_text:
        return ""
    # Normalize unicode spaces and non-breaking spaces
    text = raw_text.replace("\xa0", " ").replace("\r\n", "\n").replace("\r", "\n")
    # Collapse 3 or more newlines into 2
    text = re.sub(r"\n{3,}", "\n\n", text)
    # Strip excess horizontal whitespace on each line
    lines = [line.strip() for line in text.split("\n")]
    return "\n".join(lines).strip()


def extract_text_from_pdf(file_bytes: bytes) -> ResumeParseResult:
    """Extract text from PDF file bytes with error and scanned document detection."""
    if not fitz:
        return ResumeParseResult(
            success=False,
            file_type="pdf",
            error_message="PyMuPDF library is not installed in the environment."
        )

    if not file_bytes or len(file_bytes) == 0:
        return ResumeParseResult(
            success=False,
            file_type="pdf",
            error_message="The uploaded PDF file is empty (0 bytes)."
        )

    try:
        # Open PDF from in-memory byte buffer
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as e:
        return ResumeParseResult(
            success=False,
            file_type="pdf",
            error_message=f"Could not open PDF file. The file may be corrupted or encrypted: {str(e)}"
        )

    page_count = len(doc)
    if page_count == 0:
        doc.close()
        return ResumeParseResult(
            success=False,
            file_type="pdf",
            page_count=0,
            error_message="The uploaded PDF document contains no pages."
        )

    extracted_pages = []
    total_images = 0

    try:
        for page_idx in range(page_count):
            page = doc[page_idx]
            page_text = page.get_text() or ""
            extracted_pages.append(page_text)
            # Check for images to detect scanned documents
            image_list = page.get_images()
            total_images += len(image_list)
    except Exception as e:
        doc.close()
        return ResumeParseResult(
            success=False,
            file_type="pdf",
            page_count=page_count,
            error_message=f"Error reading pages from PDF: {str(e)}"
        )
    finally:
        doc.close()

    full_text = clean_extracted_text("\n".join(extracted_pages))
    words = full_text.split()
    word_count = len(words)

    # Scanned PDF detection: has pages and images, but very few or zero text characters
    is_scanned = (word_count < 15 and (total_images > 0 or page_count > 0 and len(full_text.strip()) == 0))
    warning = None
    if is_scanned:
        warning = (
            "Notice: Very little or no text could be extracted from this PDF. "
            "It may be a scanned image or photo. For accurate skill matching, "
            "please upload a text-based PDF or DOCX file, or paste your resume text."
        )

    return ResumeParseResult(
        success=True,
        text=full_text,
        file_type="pdf",
        page_count=page_count,
        word_count=word_count,
        is_scanned=is_scanned,
        warning_message=warning
    )


def extract_text_from_docx(file_bytes: bytes) -> ResumeParseResult:
    """Extract text from Word DOCX file bytes."""
    if not docx:
        return ResumeParseResult(
            success=False,
            file_type="docx",
            error_message="python-docx library is not installed in the environment."
        )

    if not file_bytes or len(file_bytes) == 0:
        return ResumeParseResult(
            success=False,
            file_type="docx",
            error_message="The uploaded DOCX file is empty (0 bytes)."
        )

    try:
        doc_stream = io.BytesIO(file_bytes)
        doc = docx.Document(doc_stream)
    except Exception as e:
        return ResumeParseResult(
            success=False,
            file_type="docx",
            error_message=f"Could not read DOCX document. File might be corrupted: {str(e)}"
        )

    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    
    # Also extract text from tables
    table_texts = []
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
            if row_text:
                table_texts.append(row_text)

    all_content = paragraphs + table_texts
    full_text = clean_extracted_text("\n".join(all_content))
    word_count = len(full_text.split())

    warning = None
    if word_count == 0:
        warning = "The uploaded DOCX file does not contain any readable text."

    return ResumeParseResult(
        success=True,
        text=full_text,
        file_type="docx",
        page_count=1,  # DOCX files do not have fixed page breaks by default
        word_count=word_count,
        is_scanned=False,
        warning_message=warning
    )


def extract_text_from_file(file_obj: Union[BinaryIO, bytes], filename: str) -> ResumeParseResult:
    """Universal safe dispatcher for uploaded file objects or byte buffers."""
    if hasattr(file_obj, "read"):
        file_bytes = file_obj.read()
        # Reset pointer if possible
        if hasattr(file_obj, "seek"):
            file_obj.seek(0)
    elif isinstance(file_obj, bytes):
        file_bytes = file_obj
    else:
        return ResumeParseResult(
            success=False,
            error_message="Unsupported file input type."
        )

    fn_lower = filename.lower()
    if fn_lower.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif fn_lower.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    elif fn_lower.endswith(".txt"):
        try:
            text = file_bytes.decode("utf-8", errors="replace")
            cleaned = clean_extracted_text(text)
            return ResumeParseResult(
                success=True,
                text=cleaned,
                file_type="txt",
                page_count=1,
                word_count=len(cleaned.split()),
                is_scanned=False
            )
        except Exception as e:
            return ResumeParseResult(
                success=False,
                file_type="txt",
                error_message=f"Error reading text file: {str(e)}"
            )
    else:
        return ResumeParseResult(
            success=False,
            error_message=f"Unsupported file format '{filename}'. Please upload a PDF or DOCX file."
        )
